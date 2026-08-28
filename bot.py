import argparse
import os
import re
import shutil
import subprocess
import time
from pathlib import Path
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

# ==========================================
# KONFIGURASI DAN KONSTANTA
# ==========================================
BASE = Path(__file__).resolve().parent
CHROME_PROFILE = BASE / "chrome-profile"
ARTIFACTS = BASE / "artifacts"
GENERATED_LINKS_FILE = ARTIFACTS / "generated_links.txt"

load_dotenv(BASE / ".env")

BASE_URL = os.getenv("SIPGN_BASE_URL", "https://pop-sipgn.bgn.go.id").rstrip("/")
CDP_PORT = int(os.getenv("CHROME_CDP_PORT", "9222"))
SLOW_MO = int(os.getenv("SLOW_MO", "100"))
TIMEOUT = int(os.getenv("TIMEOUT_MS", "10000"))
STOP_BEFORE_SUBMIT = os.getenv("STOP_BEFORE_SUBMIT", "true").lower() == "true"


# ==========================================
# UTILITAS DASAR
# ==========================================
def s(v):
    """Konversi nilai ke string dengan aman."""
    return "" if v is None else str(v).strip()

def num(v):
    """Konversi nilai ke string angka, membuang desimal jika berakhiran .0."""
    if v in (None, ""):
        return "0"
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else str(f)
    except Exception:
        return s(v)


# ==========================================
# MANAJEMEN BROWSER CHROME & PLAYWRIGHT
# ==========================================
def chrome_path():
    candidates = [
        os.environ.get("PROGRAMFILES", "") + r"\Google\Chrome\Application\chrome.exe",
        os.environ.get("PROGRAMFILES(X86)", "") + r"\Google\Chrome\Application\chrome.exe",
        os.environ.get("LOCALAPPDATA", "") + r"\Google\Chrome\Application\chrome.exe",
    ]
    for p in candidates:
        if p and Path(p).exists():
            return Path(p)
    found = shutil.which("chrome.exe")
    if found:
        return Path(found)
    raise FileNotFoundError("Google Chrome tidak ditemukan. Pastikan Google Chrome Desktop terpasang.")


def start_chrome():
    chrome = chrome_path()
    CHROME_PROFILE.mkdir(parents=True, exist_ok=True)

    args = [
        str(chrome),
        f"--remote-debugging-port={CDP_PORT}",
        f"--user-data-dir={CHROME_PROFILE}",
        "--no-first-run",
        "--no-default-browser-check",
        BASE_URL,
    ]

    print(f"✓ Google Chrome: {chrome}")
    print(f"✓ Chrome profile: {CHROME_PROFILE}")
    
    # Karena server sekarang dijalankan lewat CMD interaktif (Jalankan_API_SIPGN.bat), 
    # Chrome akan otomatis muncul di layar tanpa perlu workaround schtasks.
    subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def wait_cdp(timeout=30):
    import urllib.request
    endpoint = f"http://127.0.0.1:{CDP_PORT}/json/version"
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with urllib.request.urlopen(endpoint, timeout=1) as r:
                return r.read().decode("utf-8")
        except Exception:
            time.sleep(0.5)
    raise TimeoutError("Chrome CDP belum aktif setelah 30 detik.")


def connect(page_url=None):
    start_chrome()
    wait_cdp()

    pw = sync_playwright().start()
    browser = pw.chromium.connect_over_cdp(f"http://127.0.0.1:{CDP_PORT}")
    
    if not browser.contexts:
        raise RuntimeError("Tidak ada BrowserContext pada Chrome.")

    context = browser.contexts[0]
    context.set_default_timeout(TIMEOUT)

    page = context.pages[0] if context.pages else context.new_page()

    if page_url:
        page.goto(page_url, wait_until="commit", timeout=30000)
    return pw, browser, context, page


# ==========================================
# MANAJEMEN SESI & LOGIN
# ==========================================
def is_authenticated(page):
    url = page.url.lower()
    return not ("sso-sipgn.bgn.go.id/login" in url or ("/login" in url and "pop-sipgn" in url)) and "pop-sipgn.bgn.go.id" in url


def wait_manual_login(page):
    print("\n" + "=" * 65)
    print(" LOGIN MANUAL / VERIFIKASI SSO")
    print("=" * 65)
    print("Silakan selesaikan Cloudflare dan login SSO SECARA MANUAL")
    print("di Google Chrome yang baru dibuka.")
    print("Bot tidak mengklik atau menyelesaikan Cloudflare otomatis.\n")

    page.goto(f"{BASE_URL}/", wait_until="commit", timeout=30000)

    while True:
        print(f"\rHalaman saat ini: {page.url[:110]:<110}", end="")
        if is_authenticated(page):
            print("\n✓ Login terdeteksi. Melanjutkan...")
            return
        time.sleep(2)


# ==========================================
# DATA & EXCEL
# ==========================================
def read_excel(path, only_row=None):
    wb = load_workbook(path, data_only=True)
    ws = wb["DISTRIBUSI"]
    headers = [s(c.value) for c in ws[1]]
    rows = []

    for rn in range(2, ws.max_row + 1):
        row = dict(zip(headers, [ws.cell(rn, c).value for c in range(1, ws.max_column + 1)]))
        if not s(row.get("wo_id")):
            continue
        if only_row is not None and rn != only_row:
            continue
        row["_excel_row"] = rn
        rows.append(row)

    rows.sort(key=lambda r: (s(r.get("wo_id")), float(r.get("urutan_rute") or 0)))
    return rows


def validate(row):
    required = ["wo_id", "kode_wo", "urutan_rute", "tujuan_nama", "batas_waktu", "nama_driver", "plat_kendaraan"]
    errors = [x for x in required if not s(row.get(x))]
    
    for f in ["porsi_besar", "porsi_kecil", "balita", "bumil", "busui"]:
        if row.get(f) not in (None, ""):
            try:
                if float(row[f]) < 0:
                    errors.append(f"{f} negatif")
            except Exception:
                errors.append(f"{f} bukan angka")
                
    if errors:
        raise ValueError(" | ".join(errors))


def save_uuid_to_excel(excel_path, row_idx, the_uuid):
    if not excel_path:
        return
        
    try:
        wb = load_workbook(excel_path)
        ws = wb["DISTRIBUSI"]
        headers = [c.value for c in ws[1]]
        
        if "UUID" not in headers:
            col_idx = len(headers) + 1
            ws.cell(row=1, column=col_idx, value="UUID")
        else:
            col_idx = headers.index("UUID") + 1
            
        ws.cell(row=row_idx, column=col_idx, value=the_uuid)
        wb.save(excel_path)
        print(f"✓ UUID berhasil disimpan otomatis ke Excel pada baris {row_idx}")
    except PermissionError:
        print("⚠ Gagal menyimpan UUID ke Excel karena file Excel sedang dibuka/dikunci oleh program lain.")
    except Exception as e:
        print(f"⚠ Gagal menyimpan UUID ke Excel: {e}")


def record_results(row, the_uuid, excel_path):
    print(f"🔗 Link berhasil direkam: {BASE_URL}/distribution/{the_uuid}")
    
    ARTIFACTS.mkdir(exist_ok=True)
    with open(GENERATED_LINKS_FILE, "a", encoding="utf-8") as f:
        f.write(f"{row['tujuan_nama']} | {the_uuid}\n")
        
    save_uuid_to_excel(excel_path, row["_excel_row"], the_uuid)


# ==========================================
# INTERAKSI HALAMAN & INPUT
# ==========================================
def first_visible(loc):
    try:
        if loc.count() and loc.first.is_visible():
            return loc.first
    except Exception:
        pass
    return None


def click_first(locators):
    for loc in locators:
        x = first_visible(loc)
        if x:
            x.click()
            return True
    return False


def fill_first(locators, value):
    for loc in locators:
        x = first_visible(loc)
        if x:
            x.fill(s(value))
            return True
    return False


def select_wo_if_needed(page, kode_wo):
    search = page.locator('input[placeholder*="Cari kode atau nama menu" i]').first
    if not search.count():
        return

    search.fill(kode_wo)
    page.wait_for_timeout(500)
    click_first([page.get_by_text(re.compile(re.escape(kode_wo), re.I))])


def select_destination(page, name):
    search = page.locator('input[placeholder*="Cari nama KPM" i]').first
    if not search.count():
        search = page.locator('input[placeholder*="Cari nama" i]').first
    if not search.count():
        raise RuntimeError("Kolom pencarian lokasi tujuan tidak ditemukan.")

    search.fill(name)
    page.wait_for_timeout(800)

    found = click_first([
        page.get_by_text(name, exact=True),
        page.get_by_text(re.compile(re.escape(name), re.I)),
    ])
    
    if not found:
        manual_btn = page.locator('xpath=//*[contains(text(), "Tambah manual")]').first
        if manual_btn.count() and manual_btn.is_visible():
            manual_btn.click()
            manual_input = page.locator('input[placeholder*="isi manual" i]').first
            try:
                manual_input.wait_for(state="visible", timeout=3000)
                manual_input.fill(name)
            except Exception:
                search.fill(name)
        else:
            raise RuntimeError(f"Tujuan '{name}' tidak ditemukan dan tombol 'Tambah manual' tidak terlihat.")


def verify_auto_address(page):
    selectors = ['input[readonly]', 'textarea[readonly]', '[aria-label*="alamat" i]']
    for selector in selectors:
        try:
            loc = page.locator(selector)
            for i in range(loc.count()):
                el = loc.nth(i)
                if not el.is_visible():
                    continue
                try:
                    val = el.input_value()
                except Exception:
                    try:
                        val = el.inner_text()
                    except Exception:
                        val = ""
                if s(val):
                    return s(val)
        except Exception:
            pass
    return None


def fill_portions(page, row):
    page.wait_for_timeout(500)
    portions = [
        ("Porsi Besar", "porsi_besar"),
        ("Porsi Kecil", "porsi_kecil"),
        ("Balita", "balita"),
        ("Ibu Hamil", "bumil"),
        ("Ibu Menyusui", "busui"),
    ]
    
    found_portions = []
    
    for title, field in portions:
        heading = page.get_by_text(title, exact=True).first
        if heading.count() and heading.is_visible():
            found_portions.append((title, field, heading))
        else:
            heading_re = page.get_by_text(re.compile(title, re.I)).first
            if heading_re.count() and heading_re.is_visible():
                found_portions.append((title, field, heading_re))
                
    if found_portions:
        print(f"  > Porsi terdeteksi: {', '.join([p[0] for p in found_portions])}")
    else:
        print("  > Tidak ada porsi yang terdeteksi di halaman ini.")

    for title, field, heading in found_portions:
        val = row.get(field)
        card = heading.locator("xpath=ancestor::*[.//input][1]").first
        if card.count():
            inp = card.locator("input").first
            if inp.count() and inp.is_visible():
                inp.fill(num(val))
                continue
                
        following = heading.locator("xpath=following::input[1]")
        if following.count() and following.is_visible():
            following.fill(num(val))
            continue
            
        print(f"⚠ Kategori porsi '{title}' ditemukan, namun kolom inputnya tidak dapat dipetakan. Dilewati.")


def format_time(v):
    if v is None or v == "":
        return ""
    try:
        return v.strftime("%H:%M")
    except AttributeError:
        pass
    
    v_str = str(v).strip()
    parts = v_str.split(":")
    if len(parts) >= 2:
        return f"{parts[0].zfill(2)}:{parts[1].zfill(2)}"
    return v_str


def fill_route(page, row):
    if not fill_first([
        page.get_by_label(re.compile("Urutan ke", re.I)),
        page.locator('input[placeholder*="Contoh: 1" i]'),
        page.locator('xpath=//*[contains(text(), "Urutan ke-")]/following::input[1]'),
        page.locator('input[placeholder*="Urutan" i]'),
    ], num(row.get("urutan_rute"))):
        raise RuntimeError("Field Urutan rute tidak ditemukan.")

    if not fill_first([
        page.get_by_label(re.compile("Batas Waktu", re.I)),
        page.locator('xpath=//*[contains(text(), "Batas Waktu")]/following::input[1]'),
        page.locator('input[placeholder*="Batas Waktu" i]'),
    ], format_time(row.get("batas_waktu"))):
        raise RuntimeError("Field Batas Waktu tidak ditemukan.")


def fill_driver(page, row):
    driver = s(row.get("nama_driver"))
    plate = s(row.get("plat_kendaraan"))

    if not fill_first([
        page.get_by_label(re.compile("Nama Kurir|Driver|Nama.*Driver", re.I)),
        page.locator('input[placeholder*="kurir" i]'),
        page.locator('input[placeholder*="driver" i]'),
    ], driver):
        raise RuntimeError("Field nama driver/kurir tidak ditemukan.")

    page.wait_for_timeout(400)
    click_first([
        page.get_by_text(driver, exact=True),
        page.get_by_text(re.compile(re.escape(driver), re.I)),
    ])

    if not fill_first([
        page.get_by_label(re.compile("Plat", re.I)),
        page.locator('input[placeholder*="B 1234 ABC" i]'),
        page.locator('xpath=//*[contains(text(), "Plat Nomor Kendaraan")]/following::input[1]'),
        page.locator('input[placeholder*="plat" i]'),
        page.locator('input[placeholder*="nomor kendaraan" i]'),
    ], plate):
        raise RuntimeError("Field plat nomor kendaraan tidak ditemukan.")


def fill_note(page, value):
    if not s(value):
        return
    fill_first([
        page.get_by_label(re.compile("Catatan", re.I)),
        page.locator("textarea[placeholder*='catatan' i]"),
    ], value)


def upload_photos(page, row):
    paths = []
    for i in range(1, 5):
        p = s(row.get(f"foto_{i}"))
        if not p:
            continue
        path = Path(p)
        if not path.is_absolute():
            path = BASE / path
        if not path.exists():
            raise FileNotFoundError(f"Foto tidak ditemukan: {path}")
        paths.append(str(path))

    if not paths:
        return 0

    inputs = page.locator('input[type="file"]')
    if not inputs.count():
        raise RuntimeError("Input upload foto tidak ditemukan.")

    try:
        inputs.first.set_input_files(paths)
        return len(paths)
    except Exception:
        for i, p in enumerate(paths):
            if i >= inputs.count():
                break
            inputs.nth(i).set_input_files(p)
        return min(len(paths), inputs.count())


def screenshot(page, row, label):
    ARTIFACTS.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = ARTIFACTS / f"row_{row['_excel_row']}_{label}_{ts}.png"
    page.screenshot(path=str(path), full_page=True)
    return path


def inject_network_interceptor(page):
    """Memasukkan script XHR/Fetch hook untuk menangkap UUID hasil Submit."""
    page.evaluate(r"""
        window._captured_uuids = window._captured_uuids || [];
        
        if (!window._networkHooked) {
            window._networkHooked = true;
            const originalFetch = window.fetch;
            
            window.fetch = async function(...args) {
                let method = 'GET';
                if (args[0] instanceof Request) {
                    method = args[0].method;
                } else if (args[1] && args[1].method) {
                    method = args[1].method;
                }
                
                try {
                    const response = await originalFetch.apply(window, args);
                    if (['POST', 'PUT', 'PATCH'].includes(method.toUpperCase())) {
                        const clone = response.clone();
                        clone.text().then(text => {
                            const regex = /["'](?:id|data_id|distribution_id)["']\s*:\s*["']([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})["']/gi;
                            let match;
                            let foundId = false;
                            while ((match = regex.exec(text)) !== null) {
                                window._captured_uuids.push(match[1]);
                                foundId = true;
                            }
                            if (!foundId) {
                                const allUuids = text.match(/[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}/gi);
                                if (allUuids && allUuids.length === 1) {
                                    window._captured_uuids.push(allUuids[0]);
                                }
                            }
                        }).catch(e => {});
                    }
                    return response;
                } catch (err) {
                    throw err;
                }
            };
        }
    """)




def parse_datetime(value):
    import datetime
    import re
    if hasattr(value, 'strftime'):
        return value
    nums = re.findall(r'\d+', str(value))
    if len(nums) >= 5:
        if len(nums[0]) == 4:
            return datetime.datetime(int(nums[0]), int(nums[1]), int(nums[2]), int(nums[3]), int(nums[4]))
        else:
            return datetime.datetime(int(nums[2]), int(nums[1]), int(nums[0]), int(nums[3]), int(nums[4]))
    return datetime.datetime.now()

def fill_datetime_picker(page, value, label_text):
    target = parse_datetime(value)
    print(f"→ {label_text}: {target.strftime('%d/%m/%Y %H:%M')}")

    # 1. Buka Date Picker
    calendar_locator = '.grid.grid-cols-7'
    opened = False
    
    # Coba klik DOM (following-sibling)
    field = page.locator(f'xpath=//*[contains(text(), "{label_text}")]/following-sibling::*').first
    if field.count():
        try:
            field.click(force=True)
            page.wait_for_timeout(1000)
            if page.locator(calendar_locator).count():
                print("✓ Date picker dibuka (via DOM sibling)")
                opened = True
        except:
            pass

    # Coba klik DOM (following::input)
    if not opened:
        input_field = page.locator(f'xpath=//*[contains(text(), "{label_text}")]/following::input[1]')
        if input_field.count() and input_field.first.is_visible():
            try:
                input_field.first.click(force=True)
                page.wait_for_timeout(1000)
                if page.locator(calendar_locator).count():
                    print("✓ Date picker dibuka (via DOM input)")
                    opened = True
            except:
                pass
                
    # Coba klik placeholder
    if not opened:
        placeholder_input = page.locator('input[placeholder*="Pilih tanggal & waktu"]').last
        if placeholder_input.count() and placeholder_input.is_visible():
            try:
                placeholder_input.click(force=True)
                page.wait_for_timeout(1000)
                if page.locator(calendar_locator).count():
                    print("✓ Date picker dibuka (via Placeholder)")
                    opened = True
            except:
                pass

    # Fallback klik via koordinat visual jika DOM gagal / belum terbuka
    if not opened:
        label_waktu = page.locator(f'xpath=//*[contains(text(), "{label_text}") and not(*[contains(text(), "{label_text}")])]').last
        if label_waktu.count():
            box = label_waktu.bounding_box()
            if box:
                page.mouse.click(box["x"] + 20, box["y"] + box["height"] + 20)
                page.wait_for_timeout(1000)
                print("✓ Date picker dibuka (via Koordinat)")
            
    # 2. Pilih tanggal
    page.wait_for_timeout(1000)
    # Menargetkan grid kalender (Tailwind: div.grid.grid-cols-7)
    calendar = page.locator(calendar_locator).last
    
    # Kita ambil angka tanggal (kompatibel dengan semua versi Playwright)
    day_els = calendar.get_by_text(str(target.day), exact=True)
    clicked = False
    
    if day_els.count():
        # Karena grid kalender menampilkan sisa tanggal bulan sebelumnya (selalu tanggal besar > 20)
        # di awal grid, dan bulan berikutnya (selalu tanggal kecil < 15) di akhir grid.
        # Maka:
        # Jika target > 15, elemen target yg benar (di bulan saat ini) pasti berada di posisi terakhir.
        # Jika target <= 15, elemen target yg benar (di bulan saat ini) pasti berada di posisi pertama.
        if target.day > 15:
            target_el = day_els.last
        else:
            target_el = day_els.first
            
        target_el.click(force=True)
        clicked = True
        print(f"✓ Tanggal {target.day} dipilih")
        
    if not clicked:
        print(f"⚠ Tidak menemukan tanggal: {target.day} di dalam grid kalender")
        
    # 3. Pilih jam & 4. Pilih menit
    def scroll_time_column(target_val, is_hour=True):
        header_text = "JAM" if is_hour else "MENIT"
        # Pencarian teks super handal (mengabaikan spasi berlebih dan huruf besar/kecil)
        header = page.locator(f'xpath=//*[normalize-space(text())="{header_text}"] | //*[normalize-space(text())="{header_text.lower()}"] | //*[normalize-space(text())="{header_text.title()}"]').last
        
        if not header.count():
            print(f"⚠ Header {header_text} tidak ditemukan!")
            return
            
        box = header.bounding_box()
        if not box:
            print(f"⚠ Header {header_text} tidak terlihat!")
            return

        center_x = box["x"] + box["width"] / 2
        header_bottom = box["y"] + box["height"]
        max_val = 23 if is_hour else 59
        
        for attempt in range(25):
            page.wait_for_timeout(50)
            
            visible_nums = page.evaluate("""([hx, hw]) => {
                const els = Array.from(document.querySelectorAll('*')).filter(el => {
                    if (el.children.length > 0) return false;
                    const text = (el.innerText || el.textContent || '').trim();
                    return /^\\d{2}$/.test(text);
                });
                
                const headerCx = hx + hw/2;
                const columns = [];
                
                for (const el of els) {
                    const rect = el.getBoundingClientRect();
                    if (rect.width > 0 && rect.height > 0 && rect.width < 60) {
                        const style = window.getComputedStyle(el);
                        if (style.display !== 'none' && style.visibility !== 'hidden' && style.opacity !== '0') {
                            const cx = rect.x + rect.width/2;
                            let col = columns.find(c => Math.abs(c.cx - cx) < 20);
                            if (!col) {
                                col = { cx: cx, items: [] };
                                columns.push(col);
                            }
                            // Re-calculate running average for column center
                            col.cx = (col.cx * col.items.length + cx) / (col.items.length + 1);
                            col.items.push({
                                val: parseInt((el.innerText || el.textContent || '').trim(), 10),
                                y: rect.y,
                                height: rect.height,
                                cx: cx,
                                cy: rect.y + rect.height/2
                            });
                        }
                    }
                }
                
                if (columns.length === 0) return [];
                
                // Urutkan kolom berdasarkan jarak terdekat dengan titik tengah header
                columns.sort((a, b) => Math.abs(a.cx - headerCx) - Math.abs(b.cx - headerCx));
                
                // Kembalikan item pada kolom terdekat, diurutkan dari atas ke bawah
                return columns[0].items.sort((a, b) => a.y - b.y);
            }""", [box["x"], box["width"]])
            
            if not visible_nums:
                print(f"⚠ Tidak ada angka yang terlihat di bawah header {header_text}")
                break
                
            # Coba temukan target yang pas
            found = False
            for item in visible_nums:
                if item["val"] == target_val:
                    page.mouse.click(item["cx"], item["cy"])
                    found = True
                    break
                    
            if found:
                print(f"✓ {header_text} dipilih: {target_val:02d}")
                break
                
            # Jika sudah mentok 25 kali scroll dan tidak ketemu pas, pilih yang terdekat
            if attempt == 24:
                closest = min(visible_nums, key=lambda x: abs(x["val"] - target_val))
                page.mouse.click(closest["cx"], closest["cy"])
                print(f"⚠ {header_text} {target_val:02d} tidak tersedia, memilih yang terdekat: {closest['val']:02d}")
                break
                
            mid_val = visible_nums[len(visible_nums)//2]["val"]
            top_el = visible_nums[0]
            bottom_el = visible_nums[-1]
            
            up_y = (header_bottom + top_el["y"]) / 2
            gap = top_el["y"] - up_y
            down_y = bottom_el["y"] + bottom_el["height"] + gap
            
            forward_dist = (target_val - mid_val) % (max_val + 1)
            backward_dist = (mid_val - target_val) % (max_val + 1)
            
            if forward_dist <= backward_dist:
                page.mouse.click(center_x, down_y)
            else:
                page.mouse.click(center_x, up_y)
                
            page.wait_for_timeout(50)

    scroll_time_column(target.hour, is_hour=True)
    page.wait_for_timeout(300)
    scroll_time_column(target.minute, is_hour=False)
    page.wait_for_timeout(300)
    
    # 6. Verifikasi & Tutup picker
    page.keyboard.press("Escape")
    page.mouse.click(0, 0)
    page.wait_for_timeout(500)


# ==========================================
# PROSES UTAMA BOTS
# ==========================================
def process(page, row, excel_path=None):
    validate(row)

    wo_id = s(row["wo_id"])
    kode_wo = s(row["kode_wo"])
    url = f"{BASE_URL}/distribution/new?wo_id={wo_id}"
    
    print(f"\n=== EXCEL ROW {row['_excel_row']} ===")
    print(f"WO      : {kode_wo}")
    print(f"RUTE    : {row['urutan_rute']}")
    print(f"TUJUAN  : {row['tujuan_nama']}")
    print(f"DRIVER  : {row['nama_driver']}")
    print(f"PLAT    : {row['plat_kendaraan']}")

    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    select_wo_if_needed(page, kode_wo)
    print("✓ WO/Menu")

    select_destination(page, s(row["tujuan_nama"]))
    print("✓ Tujuan")

    address = verify_auto_address(page)
    print(f"✓ Alamat otomatis: {address}") if address else print("⚠ Alamat otomatis belum terbaca; tidak diisi oleh bot.")

    fill_route(page, row)
    print("✓ Urutan & batas waktu")

    fill_portions(page, row)
    print("✓ Kategori porsi")

    fill_driver(page, row)
    print("✓ Driver & kendaraan")

    fill_note(page, row.get("catatan"))
    print("✓ Catatan")

    photo_count = upload_photos(page, row)
    print(f"✓ Foto: {photo_count}")

    shot = screenshot(page, row, "before_submit")
    print(f"✓ Screenshot: {shot}")

    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        print("Silakan klik tombol Submit di Chrome.")
        
        inject_network_interceptor(page)
        
        input("Setelah notifikasi sukses muncul, tekan Enter di terminal untuk merekam kode.")
        
        captured_uuids = page.evaluate("() => window._captured_uuids || []")
        if captured_uuids:
            record_results(row, captured_uuids[-1], excel_path)
        else:
            print("⚠ Gagal menangkap kode/UUID dari network.")
            print("Pastikan submit berhasil dan tidak ada error di Chrome.")
    else:
        raise RuntimeError("Prototype tidak menyediakan submit otomatis.")


def process_berangkat(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== KEBERANGKATAN EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)
        


    waktu_val = row.get("waktu_keberangkatan")
    if waktu_val:
        fill_datetime_picker(page, waktu_val, "Waktu Keberangkatan")

    # Isi Catatan (opsional)
    catatan = row.get("catatan")
    if s(catatan):
        fill_first([
            page.locator('input[placeholder*="Catatan" i]'),
            page.locator('textarea[placeholder*="Catatan" i]'),
            page.locator('xpath=//*[contains(text(), "Catatan")]/following::input[1]')
        ], catatan)
        print("✓ Catatan (opsional) diisi")
        
    shot = screenshot(page, row, "before_berangkat")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM KEBERANGKATAN SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        print("Silakan klik tombol 'Kurir Berangkat' di Chrome.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        btn = page.get_by_role("button", name=re.compile("Kurir Berangkat", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Kurir Berangkat")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol 'Kurir Berangkat' diklik")
            page.wait_for_timeout(2000)
        else:
            raise RuntimeError("Tombol 'Kurir Berangkat' tidak ditemukan.")


def login_mode():
    pw, browser, context, page = connect()
    try:
        wait_manual_login(page)
        print("\n✓ Session Chrome sekarang sudah berada dalam kondisi login.")
        print(f"✓ Profile disimpan di:\n  {CHROME_PROFILE}\n")
        print("Anda dapat menutup Chrome atau menjalankan bot dari terminal lain.")
        print("Untuk menjaga session, jangan hapus folder chrome-profile.")
        input("Tekan Enter untuk menutup koneksi bot (Chrome tetap dapat Anda tutup sendiri)...")
    finally:
        try:
            browser.close()
        finally:
            pw.stop()


def run_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")

    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)

        for row in rows:
            try:
                process(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try:
                    screenshot(page, row, "error")
                except Exception:
                    print("⚠ (Tidak dapat mengambil screenshot error karena browser sudah tertutup)")
                
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q":
                    break
    finally:
        try:
            browser.close()
        finally:
            pw.stop()


def run_berangkat_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")

    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)

        for row in rows:
            try:
                process_berangkat(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try:
                    screenshot(page, row, "error_berangkat")
                except Exception:
                    print("⚠ (Tidak dapat mengambil screenshot error karena browser sudah tertutup)")
                
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q":
                    break
    finally:
        try:
            browser.close()
        finally:
            pw.stop()



def process_ambil(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== PENGAMBILAN EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    # 1. Jumlah Ompreng
    total_porsi = 0
    for field in ["porsi_besar", "porsi_kecil", "balita", "bumil", "busui"]:
        val = row.get(field)
        if val not in (None, ""):
            try:
                total_porsi += int(float(val))
            except Exception:
                pass
                
    # Tunggu form benar-benar siap (maksimal 10 detik)
    try:
        page.locator('input[placeholder*="25" i]').wait_for(state="visible", timeout=10000)
    except Exception:
        pass
        
    page.wait_for_timeout(500)
    
    if not fill_first([
        page.locator('input[placeholder*="25" i]')
    ], str(total_porsi)):
        print("⚠ Gagal mengisi Jumlah Ompreng.")
    else:
        print(f"✓ Jumlah Ompreng diisi: {total_porsi}")

    # 2. Nama Kurir Pengambil
    nama_driver = s(row.get("nama_driver"))
    if not fill_first([
        page.locator('input[placeholder*="Nama kurir" i]')
    ], nama_driver):
        print("⚠ Gagal mengisi Nama Kurir Pengambil.")
    else:
        print(f"✓ Nama Kurir Pengambil diisi: {nama_driver}")

    # 3. Plat Kendaraan Pengambil
    plat = s(row.get("plat_kendaraan"))
    if not fill_first([
        page.locator('input[placeholder*="9999" i]')
    ], plat):
        print("⚠ Gagal mengisi Plat Kendaraan Pengambil.")
    else:
        print(f"✓ Plat Kendaraan Pengambil diisi: {plat}")

    # 4. Waktu Dijadwalkan Pengambilan
    waktu_ambil = row.get("waktu_ambil")
    if waktu_ambil:
        fill_datetime_picker(page, waktu_ambil, "Waktu Dijadwalkan Pengambilan")
    else:
        print("⚠ Kolom 'waktu_ambil' tidak ditemukan atau kosong di Excel.")

    # 5. Foto Saat Pengambilan
    foto = row.get("bukti_foto_ambil")
    if not foto:
        print("⚠ Kolom 'bukti_foto_ambil' tidak ditemukan/kosong di Excel.")
    else:
        path = Path(s(foto))
        if not path.is_absolute():
            path = BASE / path
        if path.exists():
            inputs = page.locator('input[type="file"]')
            if inputs.count():
                inputs.first.set_input_files(str(path))
                print(f"✓ Foto saat pengambilan diunggah: {path.name}")
            else:
                print("⚠ Input upload foto tidak ditemukan.")
        else:
            print(f"⚠ Foto tidak ditemukan: {path}")
        
    shot = screenshot(page, row, "before_ambil")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM PENGAMBILAN SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        print("Silakan klik tombol 'Kurir Berangkat' di Chrome.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        btn = page.get_by_role("button", name=re.compile("Kurir Berangkat", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Kurir Berangkat")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol 'Kurir Berangkat' diklik")
            page.wait_for_timeout(2000)
        else:
            raise RuntimeError("Tombol 'Kurir Berangkat' tidak ditemukan.")


def run_ambil_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")

    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)

        for row in rows:
            try:
                process_ambil(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try:
                    screenshot(page, row, "error_ambil")
                except Exception:
                    print("⚠ (Tidak dapat mengambil screenshot error karena browser sudah tertutup)")
                
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q":
                    break
    finally:
        try:
            browser.close()
        finally:
            pw.stop()


def process_kembali(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== OMPRENG KEMBALI EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    # 1. Jumlah Ompreng (terisi otomatis)
    # 2. Nama Kurir Pengambil (terisi otomatis)
    # 3. Plat Kendaraan Pengambil (terisi otomatis)
    # 4. Waktu Dijadwalkan Pengambilan (terisi otomatis)

    # 5. Foto Kondisi Ompreng Kembali di SPPG
    foto = row.get("bukti_foto_kembali")
    if not foto:
        print("⚠ Kolom 'bukti_foto_kembali' tidak ditemukan atau kosong di Excel.")
    if foto:
        path = Path(s(foto))
        if not path.is_absolute():
            path = BASE / path
        if path.exists():
            inputs = page.locator('input[type="file"]')
            if inputs.count():
                inputs.first.set_input_files(str(path))
                print(f"✓ Foto bukti kembali diunggah: {path.name}")
            else:
                print("⚠ Input upload foto tidak ditemukan.")
        else:
            print(f"⚠ Foto tidak ditemukan: {path}")

    # 6. Waktu Ompreng Kembali di SPPG
    waktu_kembali = row.get("waktu_kembali")
    if waktu_kembali:
        fill_datetime_picker(page, waktu_kembali, "Waktu Ompreng Kembali")
    else:
        print("⚠ Kolom 'waktu_kembali' tidak ditemukan atau kosong di Excel.")

    # Isi Catatan (opsional)
    catatan = row.get("catatan")
    if s(catatan):
        fill_first([
            page.locator('input[placeholder*="Catatan" i]'),
            page.locator('textarea[placeholder*="Catatan" i]'),
            page.locator('xpath=//*[contains(text(), "Catatan")]/following::input[1]')
        ], catatan)
        print("✓ Catatan (opsional) diisi")
        
    shot = screenshot(page, row, "before_kembali")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM KEMBALI SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        print("Silakan klik tombol Submit (Simpan / Selesai) di Chrome.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        btn = page.get_by_role("button", name=re.compile("Simpan|Selesai", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Simpan"), button:has-text("Selesai")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol Submit diklik")
            page.wait_for_timeout(2000)
        else:
            print("⚠ Tombol submit tidak ditemukan otomatis.")

def run_kembali_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")

    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)

        for row in rows:
            try:
                process_kembali(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try:
                    screenshot(page, row, "error_kembali")
                except Exception:
                    print("⚠ (Tidak dapat mengambil screenshot error karena browser sudah tertutup)")
                
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q":
                    break
    finally:
        try:
            browser.close()
        finally:
            pw.stop()


def process_terima(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== TERIMA EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    # 1. Waktu Diterima di Tujuan
    waktu_diterima = row.get("waktu_diterima")
    if waktu_diterima:
        fill_datetime_picker(page, waktu_diterima, "Waktu Diterima di Tujuan")
    else:
        print("⚠ Kolom 'waktu_diterima' tidak ditemukan atau kosong di Excel.")

    # 2. Foto Bukti Terima di Tujuan
    foto = row.get("foto_bukti_terima")
    if not foto:
        print("⚠ Kolom 'foto_bukti_terima' tidak ditemukan atau kosong di Excel.")
    if foto:
        path = Path(s(foto))
        if not path.is_absolute():
            path = BASE / path
        if path.exists():
            inputs = page.locator('input[type="file"]')
            if inputs.count():
                inputs.first.set_input_files(str(path))
                print(f"✓ Foto bukti terima diunggah: {path.name}")
            else:
                print("⚠ Input upload foto tidak ditemukan.")
        else:
            print(f"⚠ Foto tidak ditemukan: {path}")

    # Isi Catatan (opsional)
    catatan = row.get("catatan")
    if s(catatan):
        fill_first([
            page.locator('input[placeholder*="Catatan" i]'),
            page.locator('textarea[placeholder*="Catatan" i]'),
            page.locator('xpath=//*[contains(text(), "Catatan")]/following::input[1]')
        ], catatan)
        print("✓ Catatan (opsional) diisi")
        
    shot = screenshot(page, row, "before_terima")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM TERIMA SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        print("Silakan klik tombol Submit di Chrome.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        # Menekan tombol submit. Menyesuaikan dengan nama tombol jika ada
        btn = page.get_by_role("button", name=re.compile("Selesai|Terima|Simpan", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Selesai"), button:has-text("Terima"), button:has-text("Simpan")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol Submit diklik")
            page.wait_for_timeout(2000)
        else:
            print("⚠ Tombol submit tidak ditemukan otomatis, abaikan jika tidak perlu.")


def run_terima_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")

    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)

        for row in rows:
            try:
                process_terima(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try:
                    screenshot(page, row, "error_terima")
                except Exception:
                    print("⚠ (Tidak dapat mengambil screenshot error karena browser sudah tertutup)")
                
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q":
                    break
    finally:
        try:
            browser.close()
        finally:
            pw.stop()

def process_mulaicuci(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== MULAI CUCI EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    # 1. Jumlah Ompreng (sama seperti proses ambil)
    total_porsi = 0
    for field in ["porsi_besar", "porsi_kecil", "balita", "bumil", "busui"]:
        val = row.get(field)
        if val not in (None, ""):
            try:
                total_porsi += int(float(val))
            except Exception:
                pass
                
    try:
        page.locator('input[placeholder*="25" i]').wait_for(state="visible", timeout=10000)
    except Exception:
        pass
        
    page.wait_for_timeout(500)
    if not fill_first([page.locator('input[placeholder*="25" i]')], str(total_porsi)):
        print("⚠ Gagal mengisi Jumlah Ompreng untuk mulai cuci.")
    else:
        print(f"✓ Jumlah Ompreng diisi: {total_porsi}")

    # 2. Waktu Dijadwalkan Mulai Pencucian
    mulai_cuci = row.get("mulai_cuci")
    if mulai_cuci:
        fill_datetime_picker(page, mulai_cuci, "Waktu Mulai Cuci")
    else:
        print("⚠ Kolom 'mulai_cuci' tidak ditemukan atau kosong di Excel.")
        
    shot = screenshot(page, row, "before_mulaicuci")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM MULAI CUCI SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        btn = page.get_by_role("button", name=re.compile("Mulai Pencucian", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Mulai Pencucian")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol 'Mulai Pencucian' diklik")
            page.wait_for_timeout(2000)
        else:
            raise RuntimeError("Tombol 'Mulai Pencucian' tidak ditemukan.")

def run_mulaicuci_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")
    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)
        for row in rows:
            try:
                process_mulaicuci(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try: screenshot(page, row, "error_mulaicuci")
                except: pass
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q": break
    finally:
        try: browser.close()
        finally: pw.stop()

def process_selesaicuci(page, row, excel_path=None):
    uuid_key = next((k for k in row.keys() if s(k).lower() == 'uuid'), None)
    uuid_val = s(row.get(uuid_key)) if uuid_key else ""
    if not uuid_val:
        kolom_tersedia = ", ".join([str(k) for k in row.keys() if k])
        raise ValueError(f"Kolom UUID kosong atau tidak ditemukan.\nKolom yang terdeteksi di Excel: {kolom_tersedia}\nPastikan baris ini benar-benar memiliki nilai UUID.")
        
    url = f"{BASE_URL}/distribution/{uuid_val}"
    
    print(f"\n=== SELESAI CUCI EXCEL ROW {row['_excel_row']} ===")
    print(f"UUID    : {uuid_val}")
    print(f"TUJUAN  : {row.get('tujuan_nama')}")
    
    page.goto(url, wait_until="commit", timeout=30000)
    page.wait_for_timeout(1500)

    if not is_authenticated(page):
        print("Session belum valid/expired.")
        wait_manual_login(page)
        page.goto(url, wait_until="commit", timeout=30000)
        page.wait_for_timeout(1200)

    # 1. Jumlah Ompreng Dicuci
    total_porsi = 0
    for field in ["porsi_besar", "porsi_kecil", "balita", "bumil", "busui"]:
        val = row.get(field)
        if val not in (None, ""):
            try:
                total_porsi += int(float(val))
            except Exception:
                pass
                
    try:
        page.locator('input[placeholder*="25" i]').wait_for(state="visible", timeout=10000)
    except Exception:
        pass
        
    page.wait_for_timeout(500)
    if not fill_first([page.locator('input[placeholder*="25" i]')], str(total_porsi)):
        print("⚠ Gagal mengisi Jumlah Ompreng Dicuci untuk selesai cuci.")
    else:
        print(f"✓ Jumlah Ompreng Dicuci: {total_porsi}")

    # 2. Waktu Selesai Cuci
    selesai_cuci = row.get("selesai_cuci")
    if selesai_cuci:
        fill_datetime_picker(page, selesai_cuci, "Waktu Selesai Cuci")
    else:
        print("⚠ Kolom 'selesai_cuci' tidak ditemukan atau kosong di Excel.")
        
    shot = screenshot(page, row, "before_selesaicuci")
    print(f"✓ Screenshot: {shot}")
    
    if STOP_BEFORE_SUBMIT:
        print("\n🟡 FORM SELESAI CUCI SUDAH DIISI. BOT BERHENTI SEBELUM SUBMIT.")
        print("Periksa data di Chrome secara manual.")
        input("Tekan Enter di terminal jika sudah selesai/submit...")
    else:
        btn = page.get_by_role("button", name=re.compile("Pencucian Selesai", re.I)).first
        if not btn.count():
            btn = page.locator('button:has-text("Pencucian Selesai")').first
            
        if btn.count() and btn.is_visible():
            btn.click()
            print("✓ Tombol 'Pencucian Selesai' diklik")
            page.wait_for_timeout(2000)
        else:
            raise RuntimeError("Tombol 'Pencucian Selesai' tidak ditemukan.")

def run_selesaicuci_mode(excel, row_number):
    rows = read_excel(excel, row_number)
    if not rows:
        raise SystemExit("Tidak ada data yang dapat diproses.")
    pw, browser, context, page = connect()
    try:
        if not is_authenticated(page):
            wait_manual_login(page)
        for row in rows:
            try:
                process_selesaicuci(page, row, excel_path=excel)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"\n❌ GAGAL row {row['_excel_row']}: {e}")
                try: screenshot(page, row, "error_selesaicuci")
                except: pass
                answer = input("Tekan Enter untuk lanjut, atau ketik q lalu Enter untuk berhenti: ")
                if answer.lower() == "q": break
    finally:
        try: browser.close()
        finally: pw.stop()

def main():
    parser = argparse.ArgumentParser(description="Hybrid SIPGN automation")
    parser.add_argument("--login", action="store_true", help="Mode login manual untuk mengamankan sesi")
    parser.add_argument("--berangkat", action="store_true", help="Jalankan bot untuk fase keberangkatan kurir")
    parser.add_argument("--ambil", action="store_true", help="Jalankan bot untuk fase pengambilan ompreng")
    parser.add_argument("--kembali", action="store_true", help="Jalankan bot untuk fase ompreng kembali")
    parser.add_argument("--terima", action="store_true", help="Jalankan bot untuk fase penerimaan tujuan")
    parser.add_argument("--mulaicuci", action="store_true", help="Jalankan bot untuk fase mulai cuci")
    parser.add_argument("--selesaicuci", action="store_true", help="Jalankan bot untuk fase selesai cuci")
    parser.add_argument("--excel", default=str(BASE / "template_bot_distribusi_sipgn_v2.xlsx"))
    parser.add_argument("--row", type=int)
    args = parser.parse_args()

    if args.login:
        login_mode()
    elif args.berangkat:
        run_berangkat_mode(args.excel, args.row)
    elif args.ambil:
        run_ambil_mode(args.excel, args.row)
    elif args.kembali:
        run_kembali_mode(args.excel, args.row)
    elif args.terima:
        run_terima_mode(args.excel, args.row)
    elif args.mulaicuci:
        run_mulaicuci_mode(args.excel, args.row)
    elif args.selesaicuci:
        run_selesaicuci_mode(args.excel, args.row)
    else:
        run_mode(args.excel, args.row)


if __name__ == "__main__":
    main()
